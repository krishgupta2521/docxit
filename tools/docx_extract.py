import json
import sys
from pathlib import Path

try:
    from docx import Document
except Exception as exc:  # pragma: no cover
    sys.stderr.write("python-docx is required. Install with: pip install python-docx\n")
    raise


def extract_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(parts)


def guess_placeholders(text: str):
    # Heuristic: extract quoted placeholders like “FIELD” or "FIELD"
    placeholders = set()
    for token in text.replace("\u201c", '"').replace("\u201d", '"').split('"'):
        token = token.strip()
        if len(token) >= 2 and token.isupper():
            placeholders.add(token)
    return sorted(placeholders)


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Missing file path"}))
        return 1
    path = Path(sys.argv[1])
    if not path.exists():
        print(json.dumps({"error": "File not found"}))
        return 1

    text = extract_text(path)
    placeholders = guess_placeholders(text)
    print(json.dumps({"text": text, "placeholders": placeholders}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
